"""
Netran AI — Embed White Theme Code Cards into Blackbook docx
Inserts the 8 generated high-resolution White-Theme code screenshots
with clean academic figure captions into Chapter 5 of the Blackbook.
"""

import shutil
import sys
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding='utf-8')

DOCX_PATH = Path(r"c:\kali\product\netran ai\doc\clg_document\NETRAN AI Blackbook (2).docx")
BACKUP_PATH = Path(r"c:\kali\product\netran ai\doc\clg_document\NETRAN AI Blackbook (2)_backup.docx")
IMG_DIR = Path(r"c:\kali\product\netran ai\doc\img_screenshort")

# Step 1: Backup original document
if not BACKUP_PATH.exists():
    shutil.copyfile(DOCX_PATH, BACKUP_PATH)
    print(f"Backed up original Blackbook to: {BACKUP_PATH.name}")
else:
    print(f"Backup already exists: {BACKUP_PATH.name}")

doc = docx.Document(DOCX_PATH)

def add_image_after_paragraph(target_p, img_filename, caption_text, width_inches=6.2):
    """Inserts a centered picture paragraph and a caption paragraph immediately after target_p."""
    img_path = IMG_DIR / img_filename
    if not img_path.exists():
        print(f"Warning: Image {img_path} not found!")
        return

    # Create new paragraph for picture
    new_p_img = doc.add_paragraph()
    new_p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    new_p_img.paragraph_format.space_before = Pt(8)
    new_p_img.paragraph_format.space_after = Pt(4)
    run_img = new_p_img.add_run()
    run_img.add_picture(str(img_path), width=Inches(width_inches))

    # Create new paragraph for caption
    new_p_cap = doc.add_paragraph()
    new_p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    new_p_cap.paragraph_format.space_before = Pt(2)
    new_p_cap.paragraph_format.space_after = Pt(12)
    run_cap = new_p_cap.add_run(caption_text)
    run_cap.font.name = "Times New Roman"
    run_cap.font.size = Pt(10)
    run_cap.font.bold = True
    run_cap.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Move elements in the XML DOM to appear right after target_p
    target_p._element.addnext(new_p_cap._element)
    target_p._element.addnext(new_p_img._element)
    print(f"  Inserted: {caption_text[:60]}...")


# Target locations in Chapter 5
insertions = [
    # 1. Section 5.2 Question Engine
    (
        "interview_questions.json + interviewer_behavior.json",
        "05_stage2_planner_analyst.png",
        "Listing 5.1: Stage 2 Question Engine — Dual Parallel Pipeline & Dynamic Whisper STT Conditioning"
    ),
    # 2. Section 5.3 TurnManager
    (
        "Echo Guard: Pre-roll buffer flushed",
        "01_stage3_turn_manager.png",
        "Listing 5.2: Stage 3 TurnManager — Level 2 State Machine & Event-Driven Echo Guard"
    ),
    # 3. Section 5.3 Silero VAD & Filter
    (
        "Web Pause Tolerance: 1200ms",
        "02_stage3_silero_vad_filter.png",
        "Listing 5.3: Stage 3 Audio DSP — Streaming IIR DC-Blocker Filter & Silero VAD ONNX Inference"
    ),
    # 4. Section 5.4 Multimodal Evaluator
    (
        "Output: report.json + evaluation_summary.md",
        "06_stage4_gemini_multimodal.png",
        "Listing 5.4: Stage 4 Evaluation — 3-Branch Gemini Multimodal Audio Evaluation with Exponential Backoff"
    ),
    # 5. Section 5.5 API Gateway
    (
        "Location: backend/api/",
        "07_api_websocket_gateway.png",
        "Listing 5.5: Web API Gateway — FastAPI Full-Duplex WebSocket Audio & Telemetry Streaming Router"
    ),
    # 6. Section 5.6 AudioWorklet
    (
        "Design System: Serene Learning Glassmorphism",
        "08_frontend_pcm_worklet.png",
        "Listing 5.6: React Web Client — AudioWorklet Native OS Decimation (48kHz → 16kHz) & Zero-Copy Transfer"
    ),
]

for match_text, img_file, caption in insertions:
    found = False
    for p in doc.paragraphs:
        if match_text in p.text:
            add_image_after_paragraph(p, img_file, caption)
            found = True
            break
    if not found:
        print(f"Could not find anchor: '{match_text}'")

# Also insert Brain and Speaker in Section 5.3 right after Listing 5.3
# Find Listing 5.2 or 5.3 anchor
target_for_brain = None
for p in doc.paragraphs:
    if "Listing 5.3: Stage 3 Audio DSP" in p.text:
        target_for_brain = p
        break

if target_for_brain:
    add_image_after_paragraph(
        target_for_brain,
        "04_stage3_speaker_flush.png",
        "Listing 5.3b: Stage 3 Speaker — Kokoro-82M Granular Sub-Chunk Streaming & Hardware Driver Flush"
    )
    add_image_after_paragraph(
        target_for_brain,
        "03_stage3_brain_llm_stream.png",
        "Listing 5.3a: Stage 3 Brain — Low-Latency LLM Streaming & Full-Duplex Barge-In Interrupt Handler"
    )

output_file = DOCX_PATH
doc.save(str(output_file))
print(f"\nSuccessfully embedded all 8 White Theme code card images into:\n{output_file}")
