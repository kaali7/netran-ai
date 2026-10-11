"""
Convert Netran AI index.md -> index.docx
Uses python-docx to produce a properly formatted Word document
matching the academic Table of Contents style.
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

# ── helpers ────────────────────────────────────────────────────────────────

def set_cell_bg(cell, hex_color: str):
    """Set a table cell background color."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_borders(cell, **kwargs):
    """Add borders to a table cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), kwargs.get(edge, "single"))
        tag.set(qn("w:sz"), "8")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), "111111")
        tcBorders.append(tag)
    tcPr.append(tcBorders)


def para_format(para, align=WD_ALIGN_PARAGRAPH.LEFT,
                space_before=0, space_after=0):
    para.alignment = align
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.space_after = Pt(space_after)


def add_run(para, text, bold=False, italic=False,
            size=12, color=None, font_name="Times New Roman"):
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = font_name
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


# ── document setup ─────────────────────────────────────────────────────────

doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3.0)
    section.right_margin  = Cm(2.5)

# Default paragraph style
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)


# ── Header block ───────────────────────────────────────────────────────────

def heading_para(text, size=12, bold=False, italic=False,
                 align=WD_ALIGN_PARAGRAPH.CENTER,
                 space_before=0, space_after=2, color=None):
    p = doc.add_paragraph()
    para_format(p, align=align, space_before=space_before, space_after=space_after)
    add_run(p, text, bold=bold, italic=italic, size=size, color=color)
    return p


heading_para("SHREE RAMCHANDRA COLLEGE OF ENGINEERING, PUNE",
             size=12, bold=True, space_before=0, space_after=3)

heading_para("Department of Computer Engineering & Artificial Intelligence",
             size=11, italic=True, space_after=6)

heading_para("NETRAN AI", size=16, bold=True, space_after=3)

heading_para("An Autonomous, Real-Time AI-Powered Technical Interviewer Platform",
             size=11, italic=True, space_after=8)

# Underlined INDEX heading
p = doc.add_paragraph()
para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=10)
run = p.add_run("INDEX (TABLE OF CONTENTS)")
run.bold = True
run.underline = True
run.font.name = "Times New Roman"
run.font.size = Pt(13)

# Divider line
doc.add_paragraph("─" * 80).paragraph_format.space_after = Pt(6)


# ── Table of contents rows ─────────────────────────────────────────────────

TOC = [
    # (sr_no,  title,                                                    page,  is_chapter)
    ("—",     "Certificate",                                             "i",      False),
    ("—",     "Acknowledgement",                                         "ii",     False),
    ("—",     "Abstract",                                                "iii",    False),
    ("—",     "List of Figures",                                         "iv",     False),
    ("—",     "List of Abbreviations",                                   "v",      False),
    # Ch1
    ("1",     "Introduction",                                            "1",      True),
    ("1.1",   "Introduction to the Topic",                               "2",      False),
    ("1.2",   "Motivation for the Project",                              "3",      False),
    ("1.3",   "Problem Definition",                                      "4",      False),
    ("1.4",   "Existing Systems",                                        "5",      False),
    ("1.5",   "Drawbacks of Existing Systems",                           "6",      False),
    ("1.6",   "Relevance of the Project",                                "7",      False),
    # Ch2
    ("2",     "Literature Survey",                                       "8",      True),
    ("2.1",   "Survey of Related Work",                                  "9",      False),
    ("2.2",   "Comparative Analysis of Existing Approaches",             "11",     False),
    ("2.3",   "Identified Research Gap",                                 "13",     False),
    # Ch3
    ("3",     "Requirements for the Proposed System",                    "14",     True),
    ("3.1",   "Introduction to Requirement Planning",                    "15",     False),
    ("3.2",   "Functional Requirements",                                 "16",     False),
    ("3.3",   "Non-Functional Requirements",                             "17",     False),
    ("3.4",   "Hardware and Software Used & Technologies Used",          "18–20",  False),
    # Ch4
    ("4",     "System Design & Architecture",                            "21",     True),
    ("4.1",   "Methodology Used",                                        "22",     False),
    ("4.2",   "System Architecture Overview",                            "23",     False),
    ("4.3",   "Four-Stage AI Pipeline Design",                           "24",     False),
    ("4.4",   "Data Flow Diagram (DFD)",                                 "25",     False),
    ("4.5",   "Entity-Relationship (ER) Diagram",                        "26",     False),
    ("4.6",   "Database Schema Design",                                  "27",     False),
    # Ch5
    ("5",     "Implementation of Proposed System",                       "28",     True),
    ("5.1",   "Stage 1 — Research Model (JD Intelligence Engine)",       "29",     False),
    ("5.2",   "Stage 2 — Question Engine (Adaptive Question Tree)",      "31",     False),
    ("5.3",   "Stage 3 — Voice-to-Voice Live Loop (Eva Interviewer)",    "33",     False),
    ("5.4",   "Stage 4 — Multimodal Evaluation & Reporting Engine",      "36",     False),
    ("5.5",   "Web API Gateway (FastAPI Backend)",                        "38",     False),
    ("5.6",   "React Web Client (Frontend SPA)",                         "40",     False),
    # Ch6
    ("6",     "Model Evaluation & Performance",                          "42",     True),
    ("6.1",   "Model Configuration & Parameters",                        "43",     False),
    ("6.2",   "STT Model Evaluation (Faster-Whisper)",                   "44",     False),
    ("6.3",   "LLM Response Quality Evaluation (Qwen 2.5)",              "45",     False),
    ("6.4",   "TTS Latency Benchmarks (Kokoro-82M)",                     "46",     False),
    ("6.5",   "Voice Activity Detection (Silero VAD) Performance",       "47",     False),
    ("6.6",   "End-to-End Pipeline Latency Analysis",                    "48",     False),
    # Ch7
    ("7",     "Testing & Results",                                       "49",     True),
    ("7.1",   "Unit Testing — Individual Stage Pipelines",               "50",     False),
    ("7.2",   "Integration Testing — Full Stack Workflow",               "51",     False),
    ("7.3",   "Code Quality Audit (Ruff, Pylint, Bandit, Radon)",        "52",     False),
    ("7.4",   "Frontend Build Verification",                             "53",     False),
    ("7.5",   "System Performance & Comparison",                         "54–56",  False),
    # Ch8
    ("8",     "Conclusion",                                              "57",     True),
    ("8.1",   "Summary of Work",                                         "58",     False),
    ("8.2",   "Limitations",                                             "58",     False),
    ("8.3",   "Future Scope & Technologies",                             "59",     False),
    ("8.4",   "References",                                              "60",     False),
]


# ── Build table ─────────────────────────────────────────────────────────────

table = doc.add_table(rows=1 + len(TOC), cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"

# Column widths
col_widths = [Cm(2.2), Cm(12.0), Cm(2.2)]
for i, width in enumerate(col_widths):
    for cell in table.columns[i].cells:
        cell.width = width

# Header row
hdr = table.rows[0]
HEADER_BG = "E8E0CC"
hdr_labels = ["Sr. No.", "Title", "Page No."]
for idx, (cell, label) in enumerate(zip(hdr.cells, hdr_labels)):
    set_cell_bg(cell, HEADER_BG)
    set_cell_borders(cell)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(label)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

# Data rows
CHAPTER_BG = "FAF7F2"
NORMAL_BG  = "FFFFFF"

for row_idx, (sr, title, page, is_chap) in enumerate(TOC):
    row = table.rows[row_idx + 1]
    bg = CHAPTER_BG if is_chap else NORMAL_BG

    # Sr. No. cell
    c0 = row.cells[0]
    set_cell_bg(c0, bg)
    set_cell_borders(c0)
    c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run(sr)
    r0.bold = is_chap
    r0.font.name = "Times New Roman"
    r0.font.size = Pt(11.5)

    # Title cell
    c1 = row.cells[1]
    set_cell_bg(c1, bg)
    set_cell_borders(c1)
    c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    # Sub-section indent
    if not is_chap and sr != "—":
        p1.paragraph_format.left_indent = Cm(0.5)
    r1 = p1.add_run(title)
    r1.bold = is_chap
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(11.5)

    # Page No. cell
    c2 = row.cells[2]
    set_cell_bg(c2, bg)
    set_cell_borders(c2)
    c2.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p2 = c2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(page)
    r2.bold = is_chap
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(11.5)


# ── Footer ─────────────────────────────────────────────────────────────────

doc.add_paragraph("")
p = doc.add_paragraph()
para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6)
add_run(p,
        "Netran AI — College Project Document  |  Academic Year 2025–26",
        italic=True, size=10, color=(100, 100, 100))


# ── Save ───────────────────────────────────────────────────────────────────

out = r"c:\kali\product\netran ai\doc\clg_document\index.docx"
doc.save(out)
print(f"Saved: {out}")
